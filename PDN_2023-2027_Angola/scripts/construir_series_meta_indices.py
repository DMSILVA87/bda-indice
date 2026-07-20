#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Constrói, a partir da execução publicada do PDN 2023-2027:

1. Séries Meta por indicador — âncoras no último dado verificado e na Meta 2027
   (publicada no PDN ou derivada da trajectória planeada), com valores intermédios
   interpolados consoante a especificação da série (linear ou geométrica).
2. Índice compósito de execução do PDN — valores observados 2022-2025 normalizados
   a 2022=100, agregados por média geométrica (domínio e global).
3. Índice compósito de execução percentual — rácios observado vs meta anual
   interpolada, ajustados à direcção, com tecto de 100%, agregados por média
   aritmética (domínio e global).

Entradas : dados/indicadores_execucao.csv, dados/metas_pdn_2027.csv
Saídas   : 03_Series_Meta_Indices_Compositos.xlsx + CSVs em dados/
"""
from __future__ import annotations

import math
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
DADOS = BASE / "dados"
SAIDA_XLSX = BASE / "03_Series_Meta_Indices_Compositos.xlsx"

ANOS_OBS = [2022, 2023, 2024, 2025]
ANO_META_FINAL = 2027

# Indicadores em que um valor mais baixo representa melhor desempenho.
MENOR_MELHOR = {
    "IMP-saude-taxa_de_incidencia_da_malaria",
    "IMP-economico-taxa_de_inflacao",
    "IMP-economico-taxa_de_desemprego",
}

# Metas 2027 publicadas no PDN com correspondência directa ao indicador de
# execução (ver dados/comparacao_execucao_metas.csv). Escala convertida quando
# necessário (carne: 788 milhares de toneladas -> toneladas).
META_2027_PUBLICADA = {
    "IMP-educacao_e_formacao_profissional-taxa_de_alfabetizacao_da_populacao_maior_que_15_anos":
        (78.0, "Meta PDN 2027 publicada (p. 55)"),
    "IMP-infraestruturas_habitacao_e_servicos_comunitarios-taxa_de_electrificacao":
        (49.0, "Meta PDN 2027 publicada (p. 104)"),
    "IMP-infraestruturas_habitacao_e_servicos_comunitarios-de_producao_total_de_energia_de_fontes_renovaveis":
        (73.0, "Meta PDN 2027 publicada (p. 104)"),
    "IMP-fomento_a_producao-toneladas_de_carne_produzidas":
        (788000.0, "Meta PDN 2027 publicada (p. 138): 788 mil t convertidas para toneladas"),
    "IMP-economico-taxa_de_crescimento_do_pib":
        (3.0, "Trajectória média PDN 2023-2027 (p. 13); não é meta anual específica"),
    "IMP-economico-taxa_de_crescimento_do_pib_nao_petrolifero":
        (4.6, "Trajectória média PDN 2023-2027 (p. 13); não é meta anual específica"),
    "IMP-economico-taxa_de_desemprego":
        (25.0, "Meta PDN 2027 publicada (p. 13)"),
}

DOMINIO_CURTO = {
    "Saúde": "Saúde",
    "Protecção Social": "Prot. Social",
    "Comunicação": "Comunicação",
    "Educação e Formação Profissional": "Educação e FP",
    "Infraestruturas, Habitação e Serviços Comunitários": "Infraestruturas",
    "Fomento à Produção": "Fomento Prod.",
    "Económico": "Económico",
}

FILL_CAB = PatternFill("solid", fgColor="1F4E5F")
FONT_CAB = Font(bold=True, color="FFFFFF")
FILL_AGG = PatternFill("solid", fgColor="DCE6F1")


# ---------------------------------------------------------------- utilitários

def gmean(valores: list[float]) -> float:
    return math.exp(sum(math.log(v) for v in valores) / len(valores))


def amean(valores: list[float]) -> float:
    return sum(valores) / len(valores)


def escolher_metodo(indicador: str, unidade: str) -> str:
    """Regra de interpolação consoante a especificação da série.

    Linear   : percentagens, taxas, rácios por habitante e stocks cumulativos
               (incrementos anuais constantes — prática recomendada para metas
               de cobertura e de infra-estrutura).
    Geométrica: níveis/quantidades não cumulativos (produção, potência),
               assumindo crescimento a taxa composta constante (CAGR).
    """
    nome = indicador.lower()
    uni = unidade or ""
    if "cumulativo" in nome:
        return "linear"
    if "%" in uni or uni.startswith("Por ") or "/100" in uni or "per capita" in uni:
        return "linear"
    if nome.startswith("taxa"):
        return "linear"
    return "geometrica"


def interpolar(v0: float, y0: int, v1: float, y1: int, metodo: str) -> dict[int, float]:
    """Série anual de y0 a y1 (inclusive). Recuo para linear se a geométrica
    não for aplicável (valores nulos/negativos)."""
    anos = range(y0, y1 + 1)
    n = y1 - y0
    if n == 0:
        return {y0: v1}
    if metodo == "geometrica" and v0 > 0 and v1 > 0:
        razao = v1 / v0
        return {t: v0 * razao ** ((t - y0) / n) for t in anos}
    return {t: v0 + (v1 - v0) * (t - y0) / n for t in anos}


def derivar_meta_2027(base: float, m25: float, metodo: str, menor_melhor: bool,
                      unidade: str) -> tuple[float, str]:
    """Meta 2027 implícita quando o PDN não publica valor equivalente.

    - Se a trajectória planeada 2022->2025 melhora o indicador, prolonga-se o
      mesmo ritmo (linear ou geométrico) até 2027.
    - Se o plano tolerava estagnação/deterioração (meta 2025 não melhor que a
      base), a meta 2025 é mantida constante até 2027 (opção conservadora).
    """
    melhora = (m25 < base) if menor_melhor else (m25 > base)
    if not melhora:
        return m25, "Derivada: Meta MINPLAN 2025 mantida constante até 2027"
    if metodo == "geometrica" and base > 0 and m25 > 0:
        m27 = base * (m25 / base) ** (5 / 3)
        origem = "Derivada: extensão geométrica da trajectória planeada 2022-2025"
    else:
        m27 = base + (m25 - base) * 5 / 3
        origem = "Derivada: extensão linear da trajectória planeada 2022-2025"
    if "%" in (unidade or "") and not menor_melhor:
        m27 = min(m27, 100.0)
    return m27, origem


def atingimento(obs: float, meta: float, menor_melhor: bool) -> float | None:
    """Percentagem de atingimento ajustada à direcção (100 = na meta)."""
    if meta is None or obs is None or meta <= 0 or obs <= 0:
        return None
    return 100.0 * (meta / obs if menor_melhor else obs / meta)


# ------------------------------------------------------------------- escrita

def escrever_tabela(ws, linha0: int, cabecalho: list[str], linhas: list[list],
                    formatos: dict[int, str] | None = None,
                    destacar: set[int] | None = None) -> int:
    """Escreve cabeçalho + linhas; devolve a linha seguinte à tabela."""
    for j, nome in enumerate(cabecalho, start=1):
        c = ws.cell(row=linha0, column=j, value=nome)
        c.fill, c.font = FILL_CAB, FONT_CAB
        c.alignment = Alignment(vertical="center", wrap_text=True)
    r = linha0 + 1
    for i, linha in enumerate(linhas):
        for j, val in enumerate(linha, start=1):
            c = ws.cell(row=r, column=j, value=val)
            if isinstance(val, float):
                if formatos and j in formatos:
                    c.number_format = formatos[j]
                else:
                    c.number_format = "#,##0" if abs(val) >= 1000 else "0.0"
            if destacar and i in destacar:
                c.fill = FILL_AGG
                c.font = Font(bold=True)
        r += 1
    return r


def ajustar_larguras(ws, larguras: dict[int, int]) -> None:
    for col, w in larguras.items():
        ws.column_dimensions[get_column_letter(col)].width = w


# ---------------------------------------------------------------------- main

def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    df = pd.read_csv(DADOS / "indicadores_execucao.csv", encoding="utf-8-sig")
    for col in ["2022_Base", "2023", "2024", "Meta_2025", "2025"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    dominios = list(dict.fromkeys(df["Domínio"]))

    # ------------------------------------------------ 1. séries meta anuais
    registos_meta = []          # metadados por indicador com meta
    metas_anuais = {}           # id -> {ano: meta}
    convergencia = {}           # id -> {ano: valor da série meta a partir do último observado}
    sem_meta = []

    for _, r in df.iterrows():
        iid, indicador, unidade = r["ID"], r["Indicador"], r["Unidade"]
        base, m25 = r["2022_Base"], r["Meta_2025"]
        menor = iid in MENOR_MELHOR
        metodo = escolher_metodo(indicador, unidade)

        obs = {t: r[str(t)] if t != 2022 else r["2022_Base"] for t in ANOS_OBS}
        obs = {t: v for t, v in obs.items() if pd.notna(v)}
        ultimo_ano = max(obs)
        ultimo_val = obs[ultimo_ano]

        if iid in META_2027_PUBLICADA:
            m27, origem = META_2027_PUBLICADA[iid]
        elif pd.notna(m25):
            m27, origem = derivar_meta_2027(base, m25, metodo, menor, unidade)
        else:
            sem_meta.append((iid, indicador))
            continue

        # trajectória planeada: base 2022 -> meta 2025 publicada -> meta 2027
        anuais = interpolar(base, 2022, m25, 2025, metodo)
        anuais.update(interpolar(m25, 2025, m27, ANO_META_FINAL, metodo))
        metas_anuais[iid] = anuais

        # série meta de convergência: último dado verificado -> meta 2027
        convergencia[iid] = interpolar(ultimo_val, ultimo_ano, m27, ANO_META_FINAL, metodo)

        atingida = (ultimo_val <= m27) if menor else (ultimo_val >= m27)
        registos_meta.append({
            "ID": iid, "Domínio": r["Domínio"], "Indicador": indicador,
            "Unidade": unidade,
            "Direcção": "Menor é melhor" if menor else "Maior é melhor",
            "Método": "Linear" if metodo == "linear" else "Geométrica (CAGR)",
            "Base_2022": float(base), "Meta_MINPLAN_2025": float(m25) if pd.notna(m25) else None,
            "Meta_2027": float(m27), "Origem_Meta_2027": origem,
            "Ano_partida": ultimo_ano, "Valor_partida": float(ultimo_val),
            "Estado": "Meta 2027 já atingida" if atingida else "Em curso",
        })

    # -------------------------------- 2. índice compósito de execução (obs.)
    normalizados = {}           # id -> {ano: índice 2022=100}
    for _, r in df.iterrows():
        iid = r["ID"]
        base = r["2022_Base"]
        menor = iid in MENOR_MELHOR
        serie = {}
        for t in ANOS_OBS:
            v = r["2022_Base"] if t == 2022 else r[str(t)]
            if pd.notna(v) and v > 0 and base > 0:
                serie[t] = 100.0 * (base / v if menor else v / base)
        normalizados[iid] = serie

    ids_painel = [i for i, s in normalizados.items() if len(s) == len(ANOS_OBS)]

    idx_dom = {d: {} for d in dominios}
    for d in dominios:
        ids_d = df.loc[df["Domínio"] == d, "ID"]
        for t in ANOS_OBS:
            vals = [normalizados[i][t] for i in ids_d if t in normalizados[i]]
            idx_dom[d][t] = gmean(vals)
    idx_global = {t: gmean([idx_dom[d][t] for d in dominios]) for t in ANOS_OBS}
    idx_painel = {}
    for t in ANOS_OBS:
        por_dom = []
        for d in dominios:
            ids_d = [i for i in df.loc[df["Domínio"] == d, "ID"] if i in ids_painel]
            por_dom.append(gmean([normalizados[i][t] for i in ids_d]))
        idx_painel[t] = gmean(por_dom)

    # ------------------------- 3. índice de execução percentual (obs vs meta)
    ANOS_ATT = [2023, 2024, 2025]
    att = {}                    # id -> {ano: % atingimento com tecto 100}
    att_livre = {}              # id -> {ano: % atingimento sem tecto}
    att_minplan25 = {}          # id -> % vs meta MINPLAN 2025 publicada (sem tecto)
    for _, r in df.iterrows():
        iid = r["ID"]
        if iid not in metas_anuais:
            continue
        menor = iid in MENOR_MELHOR
        att[iid], att_livre[iid] = {}, {}
        for t in ANOS_ATT:
            v = r[str(t)]
            if pd.isna(v):
                continue
            a = atingimento(float(v), metas_anuais[iid][t], menor)
            if a is not None:
                att_livre[iid][t] = a
                att[iid][t] = min(a, 100.0)
        if pd.notna(r["Meta_2025"]) and pd.notna(r["2025"]):
            a = atingimento(float(r["2025"]), float(r["Meta_2025"]), menor)
            if a is not None:
                att_minplan25[iid] = a

    att_dom = {d: {} for d in dominios}
    for d in dominios:
        ids_d = df.loc[df["Domínio"] == d, "ID"]
        for t in ANOS_ATT:
            vals = [att[i][t] for i in ids_d if i in att and t in att[i]]
            att_dom[d][t] = amean(vals) if vals else None
    att_global = {t: amean([att_dom[d][t] for d in dominios if att_dom[d][t] is not None])
                  for t in ANOS_ATT}
    att_global_ind = {t: amean([a[t] for a in att.values() if t in a]) for t in ANOS_ATT}

    buckets = {">100% (superada)": 0, "50-100%": 0, "<50%": 0}
    for iid, a in att_livre.items():
        if 2025 not in a:
            continue
        v = a[2025]
        if v > 100:
            buckets[">100% (superada)"] += 1
        elif v >= 50:
            buckets["50-100%"] += 1
        else:
            buckets["<50%"] += 1

    # -------------------------- 4. trajectórias das 118 metas PDN publicadas
    dfm = pd.read_csv(DADOS / "metas_pdn_2027.csv", encoding="utf-8-sig")
    for col in ["Valor_2022_Base", "Meta_2027", "Meta_2050"]:
        dfm[col] = pd.to_numeric(dfm[col], errors="coerce")
    traj_pdn, pdn_ignoradas = [], 0
    for _, r in dfm.iterrows():
        base, m27 = r["Valor_2022_Base"], r["Meta_2027"]
        if pd.isna(base) or pd.isna(m27):
            pdn_ignoradas += 1
            continue
        metodo = escolher_metodo(str(r["Indicador PDN"]), str(r["Unidade"]))
        serie = interpolar(float(base), 2022, float(m27), ANO_META_FINAL, metodo)
        traj_pdn.append((r, metodo, serie))

    # ------------------------------------------------------------- CSVs
    anos_conv = list(range(2023, ANO_META_FINAL + 1))
    rows = []
    for m in registos_meta:
        conv = convergencia[m["ID"]]
        rows.append({**{k: m[k] for k in ["ID", "Domínio", "Indicador", "Unidade",
                                          "Direcção", "Método", "Ano_partida",
                                          "Valor_partida"]},
                     **{f"S_{t}": conv.get(t) for t in anos_conv},
                     "Meta_2027": m["Meta_2027"], "Origem_Meta_2027": m["Origem_Meta_2027"],
                     "Estado": m["Estado"]})
    pd.DataFrame(rows).to_csv(DADOS / "series_meta_convergencia.csv",
                              index=False, encoding="utf-8-sig")

    rows = []
    for m in registos_meta:
        an = metas_anuais[m["ID"]]
        rows.append({**{k: m[k] for k in ["ID", "Domínio", "Indicador", "Unidade",
                                          "Direcção", "Método"]},
                     "M_2022": an[2022], **{f"M_{t}": an[t] for t in range(2023, 2028)},
                     "Origem_Meta_2027": m["Origem_Meta_2027"]})
    pd.DataFrame(rows).to_csv(DADOS / "metas_anuais_interpoladas.csv",
                              index=False, encoding="utf-8-sig")

    rows = [{"Nível": "Indicador", "ID": iid,
             "Domínio": df.loc[df["ID"] == iid, "Domínio"].iloc[0],
             "Nome": df.loc[df["ID"] == iid, "Indicador"].iloc[0],
             **{f"I_{t}": normalizados[iid].get(t) for t in ANOS_OBS}}
            for iid in normalizados]
    rows += [{"Nível": "Domínio", "ID": "", "Domínio": d, "Nome": d,
              **{f"I_{t}": idx_dom[d][t] for t in ANOS_OBS}} for d in dominios]
    rows.append({"Nível": "Global", "ID": "", "Domínio": "", "Nome": "Índice compósito de execução",
                 **{f"I_{t}": idx_global[t] for t in ANOS_OBS}})
    rows.append({"Nível": "Global", "ID": "", "Domínio": "", "Nome": "Índice (painel completo, 48 ind.)",
                 **{f"I_{t}": idx_painel[t] for t in ANOS_OBS}})
    pd.DataFrame(rows).to_csv(DADOS / "indice_composito_execucao.csv",
                              index=False, encoding="utf-8-sig")

    rows = [{"Nível": "Indicador", "ID": iid,
             "Domínio": df.loc[df["ID"] == iid, "Domínio"].iloc[0],
             "Nome": df.loc[df["ID"] == iid, "Indicador"].iloc[0],
             **{f"A_{t}": att[iid].get(t) for t in ANOS_ATT},
             "A_2025_sem_tecto": att_livre[iid].get(2025),
             "A_2025_vs_MetaMINPLAN": att_minplan25.get(iid)}
            for iid in att]
    rows += [{"Nível": "Domínio", "ID": "", "Domínio": d, "Nome": d,
              **{f"A_{t}": att_dom[d][t] for t in ANOS_ATT},
              "A_2025_sem_tecto": None, "A_2025_vs_MetaMINPLAN": None} for d in dominios]
    rows.append({"Nível": "Global", "ID": "", "Domínio": "",
                 "Nome": "Índice de execução percentual (média dos domínios)",
                 **{f"A_{t}": att_global[t] for t in ANOS_ATT},
                 "A_2025_sem_tecto": None, "A_2025_vs_MetaMINPLAN": None})
    pd.DataFrame(rows).to_csv(DADOS / "indice_execucao_percentual.csv",
                              index=False, encoding="utf-8-sig")

    rows = [{**{"ID_Meta": r["ID_Meta"], "Eixo/Política": r["Eixo/Política"],
                "Domínio/Grupo": r["Domínio/Grupo"], "Indicador": r["Indicador PDN"],
                "Unidade": r["Unidade"], "Método": "Linear" if met == "linear" else "Geométrica"},
             **{f"T_{t}": serie[t] for t in range(2022, 2028)}}
            for r, met, serie in traj_pdn]
    pd.DataFrame(rows).to_csv(DADOS / "metas_pdn_trajectorias_2022_2027.csv",
                              index=False, encoding="utf-8-sig")

    # ------------------------------------------------------------- Excel
    wb = Workbook()

    # --- Metodologia
    ws = wb.active
    ws.title = "00_Metodologia"
    texto = [
        "Séries Meta e índices compósitos do PDN 2023-2027 — metodologia",
        "",
        "Universo: 51 indicadores de execução do painel MINPLAN (Anual 2025); 46 com meta.",
        f"Sem qualquer meta publicada (excluídos das séries meta e do índice percentual): "
        f"{', '.join(n for _, n in sem_meta)}.",
        "",
        "1. SÉRIES META (folhas 01 e 02)",
        "   • Âncoras: último dado verificado (2025; 2023 nos três indicadores de professores)"
        " e Meta 2027.",
        "   • Meta 2027: publicada no PDN quando existe correspondência directa (7 casos;"
        " carne convertida de milhares de toneladas para toneladas);"
        " caso contrário, derivada da trajectória planeada 2022→Meta MINPLAN 2025:"
        " prolongamento ao mesmo ritmo se o plano previa melhoria,"
        " ou manutenção da meta 2025 se o plano tolerava estagnação/deterioração.",
        "   • Interpolação consoante a especificação da série:"
        " LINEAR para percentagens, taxas, rácios por habitante e stocks cumulativos"
        " (incrementos anuais constantes);"
        " GEOMÉTRICA (CAGR) para níveis/quantidades de produção.",
        "   • Folha 01: série de convergência do último dado verificado até à Meta 2027.",
        "   • Folha 02: metas anuais interpoladas da trajectória planeada"
        " (2022 → Meta 2025 publicada → Meta 2027), usadas no índice percentual.",
        "",
        "2. ÍNDICE COMPÓSITO DE EXECUÇÃO (folha 03) — valores observados 2022-2025",
        "   • Normalização: 2022=100; indicadores 'menor é melhor' (malária, inflação,"
        " desemprego) invertidos (base/valor).",
        "   • Agregação: média geométrica dos indicadores dentro de cada domínio e média"
        " geométrica dos 7 domínios (pesos iguais), coerente com a metodologia IGDA.",
        "   • Professores (3 indicadores) sem dados 2024-2025: o domínio Educação usa os"
        " indicadores disponíveis em cada ano; linha adicional com painel completo"
        " (48 indicadores) como robustez.",
        "",
        "3. ÍNDICE DE EXECUÇÃO PERCENTUAL (folha 04) — observado vs meta",
        "   • Atingimento anual = observado/meta interpolada (ou meta/observado quando"
        " menor é melhor), com tecto de 100% (distância-à-meta: excedentes não compensam"
        " défices).",
        "   • Agregação: média aritmética dentro do domínio e média aritmética dos domínios.",
        "   • Validação: coluna com atingimento 2025 face à meta MINPLAN publicada"
        " (sem tecto) reproduz o '% do alvo' do painel oficial nos casos comparáveis.",
        "",
        "4. TRAJECTÓRIAS DAS METAS PDN (folha 05)",
        f"   • {len(traj_pdn)} das 118 metas do PDN interpoladas de 2022 até 2027;"
        f" {pdn_ignoradas} ignoradas por falta de base 2022 ou meta numérica (marcos, etc.).",
        "",
        "Limitações: metas 2027 derivadas são implícitas e não substituem metas oficiais;"
        " as metas MINPLAN 2025 de salas de aula (I/II ciclo) são inferiores à base 2022;"
        " a meta de inflação 2025 (18%) é superior à base 2022 (14%);"
        " o crescimento do PIB usa a trajectória média 2023-2027 como referência.",
        f"Fontes: painel MINPLAN (bundle JS, Anual 2025) e PDN 2023-2027 (PDF)."
        f" Gerado por scripts/construir_series_meta_indices.py.",
    ]
    for i, t in enumerate(texto, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 130

    # --- 01 Séries Meta (convergência)
    ws = wb.create_sheet("01_Series_Meta")
    cab = (["ID", "Domínio", "Indicador", "Unidade", "Direcção", "Método",
            "Ano partida", "Valor partida"] + [f"{t}" for t in anos_conv] +
           ["Meta 2027", "Origem da Meta 2027", "Estado"])
    linhas = []
    for m in registos_meta:
        conv = convergencia[m["ID"]]
        linhas.append([m["ID"], m["Domínio"], m["Indicador"], m["Unidade"],
                       m["Direcção"], m["Método"], m["Ano_partida"], m["Valor_partida"]]
                      + [conv.get(t) for t in anos_conv]
                      + [m["Meta_2027"], m["Origem_Meta_2027"], m["Estado"]])
    escrever_tabela(ws, 1, cab, linhas)
    ws.freeze_panes = "D2"
    ajustar_larguras(ws, {1: 34, 2: 22, 3: 46, 4: 18, 5: 14, 6: 16, 7: 11, 8: 13,
                          **{8 + k: 12 for k in range(1, len(anos_conv) + 1)},
                          9 + len(anos_conv): 13, 10 + len(anos_conv): 52,
                          11 + len(anos_conv): 20})

    # --- 02 Metas anuais interpoladas
    ws = wb.create_sheet("02_Metas_Anuais")
    cab = (["ID", "Domínio", "Indicador", "Unidade", "Método"] +
           [str(t) for t in range(2022, 2028)] + ["Origem da Meta 2027"])
    linhas = []
    for m in registos_meta:
        an = metas_anuais[m["ID"]]
        linhas.append([m["ID"], m["Domínio"], m["Indicador"], m["Unidade"], m["Método"]]
                      + [an[t] for t in range(2022, 2028)] + [m["Origem_Meta_2027"]])
    escrever_tabela(ws, 1, cab, linhas)
    ws.freeze_panes = "D2"
    ajustar_larguras(ws, {1: 34, 2: 22, 3: 46, 4: 18, 5: 16,
                          **{5 + k: 13 for k in range(1, 7)}, 12: 52})

    # --- 03 Índice compósito de execução
    ws = wb.create_sheet("03_Indice_Execucao")
    fmt_idx = {k: "0.0" for k in range(4, 8)}
    cab = ["ID", "Domínio", "Indicador / agregado", "Direcção"] + [str(t) for t in ANOS_OBS]
    linhas, destaque = [], set()
    for iid, serie in normalizados.items():
        r = df.loc[df["ID"] == iid].iloc[0]
        linhas.append([iid, r["Domínio"], r["Indicador"],
                       "Menor é melhor" if iid in MENOR_MELHOR else "Maior é melhor"]
                      + [serie.get(t) for t in ANOS_OBS])
    for d in dominios:
        destaque.add(len(linhas))
        linhas.append(["", d, f"Domínio {DOMINIO_CURTO[d]} (média geométrica)", ""]
                      + [idx_dom[d][t] for t in ANOS_OBS])
    destaque.add(len(linhas))
    linhas.append(["", "", "ÍNDICE COMPÓSITO DE EXECUÇÃO (média geométrica dos domínios)", ""]
                  + [idx_global[t] for t in ANOS_OBS])
    destaque.add(len(linhas))
    linhas.append(["", "", "Índice (painel completo, 48 indicadores)", ""]
                  + [idx_painel[t] for t in ANOS_OBS])
    escrever_tabela(ws, 1, cab, linhas, formatos={5: "0.0", 6: "0.0", 7: "0.0", 8: "0.0"},
                    destacar=destaque)
    ws.freeze_panes = "D2"
    ajustar_larguras(ws, {1: 34, 2: 22, 3: 56, 4: 14, 5: 10, 6: 10, 7: 10, 8: 10})

    # --- 04 Índice de execução percentual
    ws = wb.create_sheet("04_Indice_Percentual")
    cab = (["ID", "Domínio", "Indicador / agregado", "Direcção"] +
           [f"{t} (tecto 100)" for t in ANOS_ATT] +
           ["2025 sem tecto", "2025 vs Meta MINPLAN (sem tecto)"])
    linhas, destaque = [], set()
    for iid in att:
        r = df.loc[df["ID"] == iid].iloc[0]
        linhas.append([iid, r["Domínio"], r["Indicador"],
                       "Menor é melhor" if iid in MENOR_MELHOR else "Maior é melhor"]
                      + [att[iid].get(t) for t in ANOS_ATT]
                      + [att_livre[iid].get(2025), att_minplan25.get(iid)])
    for d in dominios:
        destaque.add(len(linhas))
        linhas.append(["", d, f"Domínio {DOMINIO_CURTO[d]} (média aritmética)", ""]
                      + [att_dom[d][t] for t in ANOS_ATT] + [None, None])
    destaque.add(len(linhas))
    linhas.append(["", "", "ÍNDICE DE EXECUÇÃO PERCENTUAL (média dos domínios)", ""]
                  + [att_global[t] for t in ANOS_ATT] + [None, None])
    destaque.add(len(linhas))
    linhas.append(["", "", "Média simples dos 46 indicadores", ""]
                  + [att_global_ind[t] for t in ANOS_ATT] + [None, None])
    escrever_tabela(ws, 1, cab, linhas,
                    formatos={k: "0.0" for k in range(5, 10)}, destacar=destaque)
    ws.freeze_panes = "D2"
    ajustar_larguras(ws, {1: 34, 2: 22, 3: 56, 4: 14, 5: 14, 6: 14, 7: 14, 8: 14, 9: 18})

    r0 = len(linhas) + 4
    ws.cell(row=r0, column=1, value="Distribuição do atingimento 2025 (sem tecto, ajustado à direcção)").font = Font(bold=True)
    escrever_tabela(ws, r0 + 1, ["Escalão", "N.º de indicadores"],
                    [[k, float(v)] for k, v in buckets.items()])

    # --- 05 Trajectórias metas PDN
    ws = wb.create_sheet("05_Metas_PDN_Trajectorias")
    cab = (["ID_Meta", "Eixo/Política", "Domínio/Grupo", "Indicador PDN", "Unidade",
            "Método"] + [str(t) for t in range(2022, 2028)])
    linhas = [[r["ID_Meta"], r["Eixo/Política"], r["Domínio/Grupo"], r["Indicador PDN"],
               r["Unidade"], "Linear" if met == "linear" else "Geométrica"]
              + [serie[t] for t in range(2022, 2028)]
              for r, met, serie in traj_pdn]
    escrever_tabela(ws, 1, cab, linhas)
    ws.freeze_panes = "E2"
    ajustar_larguras(ws, {1: 14, 2: 30, 3: 24, 4: 52, 5: 22, 6: 12,
                          **{6 + k: 12 for k in range(1, 7)}})

    # --- 06 Resumo + gráficos
    ws = wb.create_sheet("06_Resumo")
    ws.cell(row=1, column=1, value="Resumo — índices compósitos do PDN 2023-2027").font = Font(bold=True, size=13)

    cab = ["Ano", "Global"] + [DOMINIO_CURTO[d] for d in dominios]
    linhas = [[str(t), idx_global[t]] + [idx_dom[d][t] for d in dominios] for t in ANOS_OBS]
    ws.cell(row=3, column=1, value="Índice compósito de execução (2022 = 100)").font = Font(bold=True)
    fim1 = escrever_tabela(ws, 4, cab, linhas, formatos={k: "0.0" for k in range(2, 10)})

    r0 = fim1 + 2
    linhas = [[str(t), att_global[t]] + [att_dom[d][t] for d in dominios] for t in ANOS_ATT]
    ws.cell(row=r0, column=1, value="Índice de execução percentual (% da meta anual, tecto 100)").font = Font(bold=True)
    fim2 = escrever_tabela(ws, r0 + 1, cab, linhas, formatos={k: "0.0" for k in range(2, 10)})
    ajustar_larguras(ws, {1: 10, **{k: 15 for k in range(2, 10)}})

    ch1 = LineChart()
    ch1.title = "Índice compósito de execução do PDN (2022=100)"
    ch1.y_axis.title = "Índice"
    ch1.x_axis.title = "Ano"
    ch1.height, ch1.width = 9, 22
    ch1.add_data(Reference(ws, min_col=2, max_col=9, min_row=4, max_row=4 + len(ANOS_OBS)),
                 titles_from_data=True)
    ch1.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(ANOS_OBS)))
    ws.add_chart(ch1, "K3")

    ch2 = LineChart()
    ch2.title = "Índice de execução percentual do PDN (% da meta, tecto 100)"
    ch2.y_axis.title = "% da meta"
    ch2.x_axis.title = "Ano"
    ch2.height, ch2.width = 9, 22
    ch2.add_data(Reference(ws, min_col=2, max_col=9, min_row=r0 + 1, max_row=r0 + 1 + len(ANOS_ATT)),
                 titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=r0 + 2, max_row=r0 + 1 + len(ANOS_ATT)))
    ws.add_chart(ch2, "K21")

    wb.save(SAIDA_XLSX)

    # ------------------------------------------------------------ consola
    print(f"OK: {SAIDA_XLSX.name}")
    print(f"Indicadores com série meta: {len(registos_meta)} | sem meta: {len(sem_meta)}")
    print(f"Metas PDN interpoladas: {len(traj_pdn)} (ignoradas: {pdn_ignoradas})")
    print("Índice compósito de execução (2022=100):",
          {t: round(idx_global[t], 1) for t in ANOS_OBS})
    print("Índice (painel completo):", {t: round(idx_painel[t], 1) for t in ANOS_OBS})
    print("Índice de execução percentual (tecto 100):",
          {t: round(att_global[t], 1) for t in ANOS_ATT})
    print("Distribuição atingimento 2025:", buckets)


if __name__ == "__main__":
    main()
