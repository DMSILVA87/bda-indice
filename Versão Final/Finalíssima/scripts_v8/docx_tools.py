"""
Geração de documentos Word a partir do modelo v7 (mantém estilos: Title, Subtitle, Author, Date,
Heading 1/2, First Paragraph, Body Text, Normal, Source Code, Table).
"""
from __future__ import annotations
import copy
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class DocBuilder:
    def __init__(self, template_path: str):
        self.doc = docx.Document(template_path)
        body = self.doc.element.body
        # remover todo o conteúdo excepto sectPr
        for child in list(body):
            if child.tag == qn("w:sectPr"):
                continue
            body.remove(child)
        self._first_after_heading = False
        self._styles = {s.style_id: s for s in self.doc.styles}
        self._by_name = {s.name: s for s in self.doc.styles}

    def _style(self, name):
        sid = name.replace(" ", "")
        return self._styles.get(sid) or self._by_name.get(name) or self._styles["Normal"]

    # ---- parágrafos ----
    def p(self, text: str, style: str = "Body Text", bold: bool = False, italic: bool = False, align=None):
        para = self.doc.add_paragraph()
        para.style = self._style(style)
        if text:
            run = para.add_run(text)
            run.bold = bold or None
            run.italic = italic or None
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return para

    def rich(self, parts, style="Body Text"):
        """parts: lista de (texto, {'bold':..,'italic':..})"""
        para = self.doc.add_paragraph(); para.style = self._style(style)
        for txt, fmt in parts:
            r = para.add_run(txt)
            r.bold = fmt.get("bold") or None
            r.italic = fmt.get("italic") or None
        return para

    def title(self, t): return self.p(t, "Title")
    def subtitle(self, t): return self.p(t, "Subtitle")
    def author(self, t): return self.p(t, "Author")
    def date(self, t): return self.p(t, "Date")
    def h1(self, t): self._first_after_heading = True; return self.p(t, "Heading 1")
    def h2(self, t): self._first_after_heading = True; return self.p(t, "Heading 2")

    def para(self, t):
        """Primeiro parágrafo após título usa 'First Paragraph', os seguintes 'Body Text'."""
        st = "First Paragraph" if self._first_after_heading else "Body Text"
        self._first_after_heading = False
        return self.p(t, st)

    def bullets(self, items, lead_bold=True):
        """Lista com marcador '•' em estilo Normal; se item for tuplo (lead, resto), o lead fica a negrito."""
        for it in items:
            para = self.doc.add_paragraph(); para.style = self._style("Normal")
            para.paragraph_format.left_indent = Pt(18)
            para.paragraph_format.first_line_indent = Pt(-12)
            if isinstance(it, tuple):
                lead, rest = it
                r = para.add_run("•  " + lead)
                r.bold = True if lead_bold else None
                para.add_run(rest)
            else:
                para.add_run("•  " + it)
        self._first_after_heading = False

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            para = self.doc.add_paragraph(); para.style = self._style("Normal")
            para.paragraph_format.left_indent = Pt(18)
            para.paragraph_format.first_line_indent = Pt(-14)
            if isinstance(it, tuple):
                lead, rest = it
                r = para.add_run(f"{i}.  " + lead); r.bold = True
                para.add_run(rest)
            else:
                para.add_run(f"{i}.  " + it)
        self._first_after_heading = False

    # ---- tabelas ----
    def table(self, header, rows, col_widths=None, font_size=9, header_fill="345A8A", zebra=True, align_num=True):
        t = self.doc.add_table(rows=1, cols=len(header))
        try:
            t.style = self.doc.styles["Table"]
        except KeyError:
            t.style = "Table Grid"
        t.autofit = True
        hdr = t.rows[0].cells
        for i, h in enumerate(header):
            hdr[i].text = ""
            run = hdr[i].paragraphs[0].add_run(str(h))
            run.bold = True; run.font.size = Pt(font_size); run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            self._shade(hdr[i], header_fill)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                txt = "" if v is None else str(v)
                run = cells[i].paragraphs[0].add_run(txt)
                run.font.size = Pt(font_size)
                if align_num and _looks_numeric(txt):
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if zebra and ri % 2 == 1:
                    self._shade(cells[i], "EEF2F7")
        if col_widths:
            for row in t.rows:
                for i, w in enumerate(col_widths):
                    if w:
                        row.cells[i].width = w
        self._set_borders(t)
        sp = self.doc.add_paragraph(); sp.style = self._style("Body Text")
        self._first_after_heading = False
        return t

    @staticmethod
    def _shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    @staticmethod
    def _set_borders(table):
        tbl = table._tbl
        tblPr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "BFC7D5")
            borders.append(el)
        tblPr.append(borders)

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)


def _looks_numeric(s: str) -> bool:
    s = s.strip().replace("−", "-").replace(" ", "").replace("%", "").replace("p.p.", "")
    if not s:
        return False
    s2 = s.replace(",", ".")
    try:
        float(s2.replace("+", ""))
        return True
    except ValueError:
        return False


def fmt(x, nd=1, pct=False, sign=False):
    """Formato PT: vírgula decimal; None -> '–'."""
    if x is None or x == "":
        return "–"
    if isinstance(x, str):
        return x
    s = f"{x:+.{nd}f}" if sign else f"{x:.{nd}f}"
    s = s.replace(".", ",").replace("-", "−")
    if pct:
        s += "%"
    return s


def fmt_int(x):
    if x is None:
        return "–"
    return f"{int(round(x)):,}".replace(",", " ")
