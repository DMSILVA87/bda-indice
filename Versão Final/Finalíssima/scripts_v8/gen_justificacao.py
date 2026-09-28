"""
Gera IGDA_BDA_Justificacao_Minimos_Maximos_v8.docx a partir da Base_Potencial do construtor v8.
Corrige o erro de texto da v7 (escalas 0-100 descritas como "0 e 1") e actualiza séries/metas.
Uso: python3 gen_justificacao.py <construtor_v8.xlsx> <saida.docx> [<v8_results.json>]
"""
from __future__ import annotations
import sys, os, json
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import openpyxl
from docx.shared import Cm
from docx_tools import DocBuilder, fmt
import v8_plan as P

REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
FILES = os.path.join(REPO, "Versão Final/Finalíssima/files")
DIMS = ["1. Governança e Estado de Direito", "2. Estabilidade Macroeconómica", "3. Capital Humano", "4. Inclusão Social e Protecção",
        "5. Infraestruturas e Serviços", "6. Mercado de Trabalho", "7. Segurança Alimentar e Saúde", "8. Diversificação Produtiva e Sector Privado",
        "9. Ambiente, Clima e Resiliência", "10. Demografia, Território e Urbanização", "11. Transformação Digital e Inovação"]

REVISED_V7 = {"INF024", "LAB003", "LAB025"}
META_PHRASE = {
    "pdn": "Meta alinhada com o PDN 2023-2027 (v8).",
    "transposta": "Meta do PDN 2023-2027 transposta à base da série (v8; ver coluna 'Origem da Meta 2027').",
    "convertida": "Meta do PDN 2023-2027 convertida de escala/unidade (v8).",
    "minplan": "Meta anual 2025 do Balanço do PDN (MINPLAN) usada como proxy da meta 2027 (v8).",
    "operacional": "Meta operacional da v7 mantida (sem correspondência directa no PDN; ver coluna 'Origem da Meta 2027').",
}


def num_pt(x):
    if x is None or x == "":
        return "–"
    if isinstance(x, str):
        return x
    if abs(x) >= 1000 and float(x).is_integer():
        return f"{int(x):,}".replace(",", " ")
    if float(x).is_integer():
        return str(int(x))
    return f"{x:.4g}".replace(".", ",") if abs(x) < 1000 else f"{x:,.1f}".replace(",", " ").replace(".", ",")


def justify(row, cols, years, meta_cat):
    lo, hi = row[cols["Mínimo"]], row[cols["Máximo"]]
    unit = str(row[cols["Unidade"]] or "")
    sent = str(row[cols["Sentido"]] or "")
    meta = row[cols["Meta 2027"]]
    vals = [row[cols[y]] for y in years if isinstance(row[cols[y]], (int, float))]
    u = unit.lower()
    sl, sh = num_pt(lo), num_pt(hi)
    # tipo de escala
    if "-2.5" in unit or "2,5" in unit:
        base = f"Os limites {sl} e {sh} reproduzem a escala metodológica oficial dos Worldwide Governance Indicators, em que valores mais altos representam melhor governação."
    elif unit.strip() in ("1-6",):
        base = "Os limites 1 e 6 seguem a escala oficial CPIA do Banco Mundial, em que 1 é desempenho fraco e 6 é desempenho forte."
    elif unit.strip() in ("1-5",):
        base = "Os limites 1 e 5 reproduzem a escala oficial do índice (1 = pior, 5 = melhor)."
    elif unit.strip() in ("0-1",) or (lo == 0 and hi == 1):
        base = "Os limites 0 e 1 seguem a escala normalizada do índice/proporção: 0 representa o mínimo e 1 o máximo metodológico da escala."
    elif unit.strip().startswith("0-100") or unit.strip() in ("0-100", "0-100 (maior=mais livre)", "0-120", "0-10"):
        base = f"Os limites {sl} e {sh} correspondem à escala oficial publicada pela fonte (score), preservando a leitura directa de pior a melhor desempenho."
    elif "%" in unit and lo == 0 and hi == 100:
        base = "Os limites 0 e 100 reflectem o intervalo natural de uma percentagem/proporção: não pode ser inferior a zero nem superior à totalidade do universo medido."
    elif sent == "+/-":
        base = f"O intervalo {sl} a {sh} delimita uma zona operacional para uma variável de equilíbrio ou composição; fora dessa zona o valor deixa de acrescentar sinal comparável ao índice."
    elif isinstance(lo, (int, float)) and lo < 0:
        base = f"Como o indicador pode assumir valores negativos e positivos, o intervalo {sl} a {sh} capta deterioração e melhoria sem tornar a escala excessivamente sensível a choques."
    elif "ano" in u and "%" not in u:
        base = f"O intervalo {sl} a {sh} usa limites substantivos plausíveis para variáveis medidas em anos; evita pisos irreais e mantém a normalização comparável."
    elif sent == "-":
        if any(k in u for k in ("por 1", "por 100", "‰", "número", "nv")):
            base = f"O mínimo {sl} corresponde ao melhor caso teórico ou incidência nula. O máximo {sh} é um limite adverso plausível para transformar taxas de incidência em pontuação sem premiar diferenças para além de patamares críticos."
        else:
            base = f"O mínimo {sl} representa ausência do problema ou risco medido. O máximo {sh} é um tecto operacional de penalização antes de truncagem na escala normalizada."
    else:
        if lo == 0 and any(k in u for k in ("usd", "kz", "milh", "número", "teu", "km", "kwh", "ton")):
            base = f"O mínimo 0 é o piso natural de uma contagem ou variável monetária não negativa. O máximo {sh} é um tecto operacional calibrado pela ordem de grandeza observada, metas e benchmarks, evitando que valores excepcionais dominem a normalização."
        elif lo == 0:
            base = f"O mínimo 0 representa ausência ou valor nulo do fenómeno positivo. O máximo {sh} é um benchmark operacional compatível com metas, pares internacionais e a escala histórica observada."
        else:
            base = f"O intervalo {sl} a {sh} é uma faixa operacional para uma variável monetária/escala de grandeza que não parte realisticamente de zero no caso observado."
    parts = [base]
    if row[cols["ID"]] in REVISED_V7:
        parts.append("Fronteiras revistas na v7: as anteriores derivavam da amplitude histórica observada, o que comprimia a escala e atribuía, por construção, o score máximo ao melhor ano da série.")
    if isinstance(meta, (int, float)):
        inside = (lo <= meta <= hi) if isinstance(lo, (int, float)) and isinstance(hi, (int, float)) else True
        parts.append(f"A meta 2027 ({num_pt(meta)}) fica dentro da faixa, permitindo interpretar progresso sem saturar a escala antes do objectivo." if inside else f"A meta 2027 ({num_pt(meta)}) está fora da faixa: rever limites ou meta.")
        cat = meta_cat.get(row[cols["ID"]], "operacional")
        parts.append(META_PHRASE.get(cat, META_PHRASE["operacional"]))
    if vals:
        parts.append(f"A série observada em 2015-2025 vai de {num_pt(min(vals))} a {num_pt(max(vals))}.")
    else:
        parts.append("A linha ainda não tem série anual consolidada em 2015-2025; os limites são parâmetros metodológicos para futura incorporação.")
    return " ".join(parts)


def build(xlsx_path, template_path, out_path, results_path, version="v8", date="Setembro de 2026"):
    R = json.load(open(results_path))
    meta_cat = R["meta_cat"]
    new_ids = [e["id"] for e in R["log"] if e["tipo"] == "novo indicador"]
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["02_Base_Potencial"]
    hdr = [c.value for c in ws[4]]
    cols = {h: i for i, h in enumerate(hdr) if h is not None}
    years = list(range(2015, 2026))
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    b = DocBuilder(template_path)
    b.title(f"IGDA-BDA — Justificação dos limites mínimo e máximo ({version})")
    b.subtitle(f"Base Potencial do construtor do índice (IGDA_BDA_Construtor_{version}.xlsx)")
    b.author("Banco de Desenvolvimento de Angola · Gabinete de Planeamento e Controlo")
    b.date(date)
    b.h1("Nota metodológica")
    b.para("Este documento justifica os limites de mínimo e máximo usados na folha 02_Base_Potencial. Esses limites são parâmetros de normalização min-max: não são previsões, nem necessariamente limites físicos absolutos. "
           "Quando a fonte tem escala oficial, o intervalo replica essa escala. Quando a variável é uma percentagem ou proporção, usa-se o limite natural sempre que aplicável. Quando a variável não tem tecto natural, "
           "o limite é operacional: serve para preservar comparabilidade, evitar que outliers dominem o índice e manter a meta 2027 dentro da faixa sempre que possível.")
    n_in = sum(1 for r in rows if all(isinstance(r[cols[k]], (int, float)) for k in ("Mínimo", "Máximo")) and all((r[cols[y]] is None) or (r[cols["Mínimo"]] <= r[cols[y]] <= r[cols["Máximo"]]) for y in years))
    b.para(f"Validação: todos os {len(rows)} indicadores da Base Potencial têm mínimo e máximo definidos e {n_in} têm todas as observações de 2015-2025 dentro dos respectivos intervalos" + ("." if n_in == len(rows) else f" (os restantes {len(rows) - n_in} são assinalados na tabela)."))
    numeros = {7: "sete", 6: "seis", 8: "oito"}.get(len(new_ids), str(len(new_ids)))
    b.para(f"Revisão {version}: nenhuma fronteira dos 41 indicadores do índice foi alterada face à v7. As alterações desta versão incidem sobre (i) a coluna Meta 2027, agora alinhada com o PDN 2023-2027 sempre que existe correspondência — "
           "directa, transposta à base da série ou convertida de escala — com a origem e a categoria de cada meta documentadas na nova coluna 'Origem da Meta 2027' da Base_Potencial; (ii) as séries actualizadas com as publicações mais recentes; "
           f"e (iii) {numeros} novos candidatos ({', '.join(new_ids)}), cujos limites seguem as regras abaixo. "
           "Foi ainda corrigido um erro de redacção da v7, que descrevia várias escalas 0-100 como '0 e 1'.")
    b.table(["Critério", "Justificação geral"], [
        ["Escalas oficiais", "WGI (−2,5 a 2,5), CPIA (1 a 6), UHC e scores 0-100 mantêm a escala publicada pela fonte."],
        ["Percentagens/proporções", "Quando a percentagem é naturalmente limitada, usa-se 0-100; quando é razão macroeconómica, usa-se faixa operacional plausível."],
        ["Indicadores negativos", "O mínimo tende a representar ausência ou baixa incidência do problema; o máximo representa patamar adverso para truncagem/penalização."],
        ["Indicadores positivos", "O mínimo representa ausência, piso técnico ou patamar baixo; o máximo representa benchmark elevado ou fronteira de desenvolvimento."],
        ["Variáveis monetárias/contagens", "O mínimo é normalmente zero; o máximo é tecto operacional baseado em ordem de grandeza, meta, série observada e comparabilidade."],
        ["Fronteiras de amplitude histórica", "Não são usadas em nenhum indicador que integre o índice (revisão v7). Candidatos fora do índice que ainda as usam mantêm os parâmetros para eventual revisão quando forem incorporados."],
        ["Variáveis +/−", "São variáveis de equilíbrio ou composição; o intervalo delimita a zona de leitura útil para normalização."],
        ["Metas 2027 (v8)", "Alinhadas com o PDN 2023-2027 por cinco vias: valor directo (base do PDN coincide com a série); transposição à base da série quando a base 2022 do PDN difere (variação aditiva para níveis/proporções, relativa para mortalidade, desemprego e dívida); "
                            "conversão de escala (percentis WGI → estimativas; % do PIB não petrolífero → % do PIB); meta anual 2025 do MINPLAN como proxy; meta operacional da v7 mantida quando não há correspondência. Sem meta quando o conceito do PDN difere do da série. A coluna Meta 2027 não entra em nenhuma fórmula."],
    ], col_widths=[Cm(4.5), Cm(20)])
    b.h1("Justificação linha a linha")
    for dim in DIMS:
        sub = [(i, r) for i, r in enumerate(rows) if r[cols["Dimensão"]] == dim]
        if not sub:
            continue
        b.h2(dim)
        table_rows = []
        for i, r in sub:
            table_rows.append([i + 5, r[cols["ID"]], r[cols["Indicador"]], f"{r[cols['Unidade']]} / {r[cols['Sentido']]}", num_pt(r[cols["Mínimo"]]), num_pt(r[cols["Máximo"]]), justify(r, cols, years, meta_cat)])
        b.table(["Linha", "ID", "Indicador", "Unidade / sentido", "Mín.", "Máx.", "Justificação"], table_rows, col_widths=[Cm(1.2), Cm(1.7), Cm(4.6), Cm(2.6), Cm(1.4), Cm(1.6), Cm(11.5)], font_size=8)
    b.para(f"Total de linhas documentadas: {len(rows)}. Fonte: folha 02_Base_Potencial de IGDA_BDA_Construtor_{version}.xlsx.")
    b.save(out_path)
    return len(rows)


if __name__ == "__main__":
    results = sys.argv[3] if len(sys.argv) > 3 else os.path.join(HERE, "v8_results.json")
    n = build(sys.argv[1], os.path.join(FILES, "IGDA_BDA_Justificacao_Minimos_Maximos_v7.docx"), sys.argv[2], results)
    print("rows", n)
