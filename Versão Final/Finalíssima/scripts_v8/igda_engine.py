"""
Motor de cálculo em Python que replica a cadeia de fórmulas do IGDA_BDA_Construtor (v7/v8).

Replica exactamente:
  02_Base_Potencial: scoring (AB..AK), elegibilidade (AL), melhor do grupo (AN), rank (AO),
                     selecção (AQ) e slot (AR)
  03_Base_Indice:    indicadores seleccionados por dimensão (até 10 slots)
  04_Dados_Ajustados: imputação (interpolação linear / carry-forward / carry-backward)
  05_Flags_Qualidade
  06_Normalizacao:   min-max com sentido
  07_Subindices:     média aritmética ponderada por dimensão (cobertura mínima) + IGDA geométrico
  11_Escala_Comum:   2015 = 100

Uso:
    from igda_engine import load_inputs_from_workbook, Model
    inp = load_inputs_from_workbook(path)
    m = Model(inp); m.run(); m.subindices, m.igda
"""
from __future__ import annotations
import math
from dataclasses import dataclass, field
from typing import Any

import openpyxl

YEARS = list(range(2015, 2026))
NY = len(YEARS)

DIM_ORDER = [
    "1. Governança e Estado de Direito",
    "2. Estabilidade Macroeconómica",
    "3. Capital Humano",
    "4. Inclusão Social e Protecção",
    "5. Infraestruturas e Serviços",
    "6. Mercado de Trabalho",
    "7. Segurança Alimentar e Saúde",
    "8. Diversificação Produtiva e Sector Privado",
    "9. Ambiente, Clima e Resiliência",
    "10. Demografia, Território e Urbanização",
    "11. Transformação Digital e Inovação",
]
DIM_SHORT = ["Governança", "Macroeconomia", "Capital Humano", "Inclusão Social", "Infraestruturas",
             "Mercado Trabalho", "Saúde/Alimentar", "Diversificação", "Ambiente", "Demografia", "Digital/Inovação"]


def canon_dim(name) -> str:
    """Normaliza a grafia do nome da dimensão (a v7 usava 'Setor Privado'; a v8 uniformiza para 'Sector Privado')."""
    s = str(name or "")
    return s.replace("Setor Privado", "Sector Privado")


def _num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def excel_round(x: float, nd: int) -> float:
    """ROUND do Excel (half away from zero)."""
    if x is None:
        return None
    m = 10 ** nd
    if x >= 0:
        return math.floor(x * m + 0.5) / m
    return -math.floor(-x * m + 0.5) / m


@dataclass
class Row:
    excel_row: int
    id: str
    dim: str
    subtema: str
    indicador: str
    unidade: str
    sentido: str
    minimo: Any
    maximo: Any
    meta: Any
    peso: Any
    fonte: str
    codigo: str
    origem: str
    nacional: Any
    grupo: str
    prioridade: str
    values: list  # 11 values (None or float)
    ajuste: str = "Auto"
    nota: str = ""
    # computed
    nobs: int = 0
    cobertura: float = 0.0
    primeiro: Any = ""
    ultimo: Any = ""
    continuidade: Any = ""
    recencia: Any = ""
    s_rel: float = 0.0
    s_comp: float = 0.0
    s_cons: float = 0.0
    score: float = 0.0
    elegivel: int = 0
    motivo: str = ""
    melhor_grupo: int = 0
    rank: Any = ""
    seleccionado: int = 0
    slot: Any = ""


@dataclass
class Params:
    ano_inicial: int = 2015
    ano_final: int = 2025
    esc_min: float = 0.0
    esc_max: float = 100.0
    cobertura_min: float = 0.8
    recencia_min: int = 2023
    gap_interp: int = 2
    carry_fwd: int = 2
    carry_bwd: int = 2
    cobertura_sub: float = 0.75
    min_subindices: int = 4
    modo: str = "Automática"
    w_rel: float = 0.4
    w_comp: float = 0.35
    w_cons: float = 0.25
    bonus_nacional: float = 12.0
    dims_incluir: dict = field(default_factory=dict)   # dim -> 1/0
    dims_peso: dict = field(default_factory=dict)      # dim -> peso
    dims_alvo: dict = field(default_factory=dict)      # dim -> N alvo


@dataclass
class Inputs:
    rows: list
    params: Params


def load_inputs_from_workbook(path: str, sheet_base: str = "02_Base_Potencial", sheet_crit: str = "01_Criterios") -> Inputs:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_base]
    hdr = [c.value for c in ws[4]]
    col = {h: i for i, h in enumerate(hdr)}
    rows = []
    for r in ws.iter_rows(min_row=5, max_row=ws.max_row):
        vals = [c.value for c in r]
        if vals[col["ID"]] in (None, ""):
            continue
        yv = []
        for y in YEARS:
            v = vals[col[y]] if y in col else None
            yv.append(float(v) if _num(v) else None)
        rows.append(Row(
            excel_row=r[0].row,
            id=str(vals[col["ID"]]), dim=canon_dim(vals[col["Dimensão"]]), subtema=str(vals[col["Subtema"]] or ""),
            indicador=str(vals[col["Indicador"]] or ""), unidade=str(vals[col["Unidade"]] or ""),
            sentido=str(vals[col["Sentido"]] or ""), minimo=vals[col["Mínimo"]], maximo=vals[col["Máximo"]],
            meta=vals[col["Meta 2027"]], peso=vals[col["Peso"]], fonte=str(vals[col["Fonte"]] or ""),
            codigo=str(vals[col["Código/API"]] or ""), origem=str(vals[col["Origem"]] or ""),
            nacional=vals[col["Fonte nacional"]], grupo=str(vals[col["Grupo (anti-duplicação)"]] or ""),
            prioridade=str(vals[col["Prioridade"]] or ""), values=yv,
            ajuste=str(vals[col["Ajuste manual"]] or "Auto"), nota=str(vals[col["Nota"]] or ""),
        ))
    wc = wb[sheet_crit]
    p = Params()
    p.ano_inicial = int(wc["C7"].value); p.ano_final = int(wc["C8"].value)
    p.esc_min = float(wc["C9"].value); p.esc_max = float(wc["C10"].value)
    p.cobertura_min = float(wc["C11"].value); p.recencia_min = int(wc["C12"].value)
    p.gap_interp = int(wc["C13"].value); p.carry_fwd = int(wc["C14"].value); p.carry_bwd = int(wc["C15"].value)
    p.cobertura_sub = float(wc["C16"].value); p.min_subindices = int(wc["C17"].value)
    p.modo = str(wc["C18"].value)
    p.w_rel = float(wc["C19"].value); p.w_comp = float(wc["C20"].value); p.w_cons = float(wc["C21"].value)
    p.bonus_nacional = float(wc["C22"].value)
    for r in range(26, 37):
        d = canon_dim(wc[f"B{r}"].value) if wc[f"B{r}"].value else None
        if d:
            p.dims_incluir[d] = wc[f"C{r}"].value
            p.dims_peso[d] = wc[f"D{r}"].value
            p.dims_alvo[d] = wc[f"E{r}"].value
    return Inputs(rows=rows, params=p)


class Model:
    def __init__(self, inp: Inputs):
        self.rows = inp.rows
        self.p = inp.params
        self.base_indice = {}      # dim -> list of Row (por slot)
        self.adjusted = {}         # (dim, slot) -> list[11]
        self.flags = {}            # (dim, slot) -> list[11]
        self.scores = {}           # (dim, slot) -> list[11]
        self.subindices = {}       # dim -> list[11] (None = "")
        self.igda = [None] * NY
        self.escala = {}           # dim -> list[11]

    # ---------- 02_Base_Potencial ----------
    def score_rows(self):
        p = self.p
        for r in self.rows:
            obs = [v for v in r.values if v is not None]
            r.nobs = len(obs)
            r.cobertura = r.nobs / NY
            if r.nobs:
                idx = [i for i, v in enumerate(r.values) if v is not None]
                r.primeiro = YEARS[idx[0]]; r.ultimo = YEARS[idx[-1]]
                r.continuidade = r.nobs / (r.ultimo - r.primeiro + 1)
                r.recencia = p.ano_final - r.ultimo
            else:
                r.primeiro = r.ultimo = r.continuidade = r.recencia = ""
            pr = r.prioridade
            r.s_rel = 100 if pr == "Em uso" else 85 if pr == "Alta" else 55 if pr == "Média" else 30
            r.s_comp = excel_round(100 * r.cobertura, 1)
            r.s_cons = 0 if r.nobs == 0 else excel_round(max(0, 100 * r.continuidade - 10 * r.recencia), 1)
            nat = r.nacional if _num(r.nacional) else 0
            r.score = excel_round((p.w_rel * r.s_rel + p.w_comp * r.s_comp + p.w_cons * r.s_cons) / (p.w_rel + p.w_comp + p.w_cons) + nat * p.bonus_nacional, 1)
            ok_sent = r.sentido in ("+", "-")
            r.elegivel = 1 if (r.nobs >= 3 and r.cobertura >= p.cobertura_min and (r.ultimo if _num(r.ultimo) else 0) >= p.recencia_min
                               and ok_sent and _num(r.minimo) and _num(r.maximo)) else 0
            if r.elegivel:
                r.motivo = ""
            else:
                m = ""
                if r.nobs == 0: m += "sem dados no período; "
                elif r.nobs < 3: m += "série demasiado curta; "
                if r.nobs > 0 and r.cobertura < p.cobertura_min: m += "cobertura insuficiente; "
                if r.nobs > 0 and (r.ultimo if _num(r.ultimo) else 0) < p.recencia_min: m += "desactualizado; "
                if not ok_sent: m += "sentido ambíguo; "
                if not (_num(r.minimo) and _num(r.maximo)): m += "sem limites min/max; "
                r.motivo = m.strip()
        # melhor do grupo
        for r in self.rows:
            if r.elegivel == 0 or r.ajuste == "Excluir":
                r.melhor_grupo = 0
                continue
            better = 0
            for o in self.rows:
                if o.grupo == r.grupo and o.elegivel == 1 and o.ajuste != "Excluir":
                    if o.score > r.score or (o.score == r.score and o.excel_row < r.excel_row):
                        better += 1
            r.melhor_grupo = 1 if better == 0 else 0
        # rank na dimensão
        for r in self.rows:
            if r.melhor_grupo == 0:
                r.rank = ""
                continue
            better = 0
            for o in self.rows:
                if o.dim == r.dim and o.melhor_grupo == 1:
                    if o.score > r.score or (o.score == r.score and o.excel_row < r.excel_row):
                        better += 1
            r.rank = 1 + better
        # selecção
        for r in self.rows:
            inc = self.p.dims_incluir.get(r.dim, 0)
            if inc != 1:
                r.seleccionado = 0
            elif r.ajuste == "Excluir":
                r.seleccionado = 0
            elif r.ajuste == "Incluir":
                r.seleccionado = 1 if (r.sentido in ("+", "-") and _num(r.minimo) and _num(r.maximo) and r.nobs > 0) else 0
            elif self.p.modo == "Automática":
                alvo = self.p.dims_alvo.get(r.dim, 0)
                r.seleccionado = 1 if (r.melhor_grupo == 1 and r.rank != "" and r.rank <= alvo) else 0
            else:
                r.seleccionado = 0
        # slot
        for r in self.rows:
            if r.seleccionado == 0:
                r.slot = ""
                continue
            same = [o for o in self.rows if o.dim == r.dim and o.seleccionado == 1]
            if r.rank != "":
                r.slot = 1 + sum(1 for o in same if o.rank != "" and o.rank < r.rank)
            else:
                r.slot = 1 + sum(1 for o in same if o.rank != "") + sum(1 for o in same if o.rank == "" and o.excel_row < r.excel_row)

    # ---------- 03 / 04 / 05 / 06 ----------
    def build_index(self):
        p = self.p
        for dim in DIM_ORDER:
            slots = {}
            for r in self.rows:
                if r.dim == dim and r.seleccionado == 1 and r.slot != "":
                    slots[r.slot] = r
            self.base_indice[dim] = [slots.get(s) for s in range(1, 11)]
            for s in range(1, 11):
                r = slots.get(s)
                if r is None:
                    continue
                vals = r.values
                adj, flg, sc = [], [], []
                for t in range(1, NY + 1):  # t = 1..11
                    v = vals[t - 1]
                    if v is not None:
                        adj.append(v); flg.append("OBS")
                        continue
                    prev_idx = max([i + 1 for i in range(0, t) if vals[i] is not None], default=None)
                    next_idx = min([i + 1 for i in range(t - 1, NY) if vals[i] is not None], default=None)
                    if prev_idx is not None and next_idx is not None and (next_idx - prev_idx - 1) <= p.gap_interp:
                        pv, nv = vals[prev_idx - 1], vals[next_idx - 1]
                        adj.append(pv + (nv - pv) * (t - prev_idx) / (next_idx - prev_idx)); flg.append("EST_INTERP")
                    elif prev_idx is not None and (t - prev_idx) <= p.carry_fwd:
                        adj.append(vals[prev_idx - 1]); flg.append("EST_CARRY_FWD")
                    elif next_idx is not None and (next_idx - t) <= p.carry_bwd:
                        adj.append(vals[next_idx - 1]); flg.append("EST_CARRY_BWD")
                    else:
                        adj.append(None); flg.append("SEM_DADO")
                lo, hi = float(r.minimo), float(r.maximo)
                for a in adj:
                    if a is None:
                        sc.append(None); continue
                    if r.sentido == "+":
                        x = p.esc_min + (p.esc_max - p.esc_min) * (a - lo) / (hi - lo)
                    else:
                        x = p.esc_min + (p.esc_max - p.esc_min) * (hi - a) / (hi - lo)
                    sc.append(min(p.esc_max, max(p.esc_min, x)))
                self.adjusted[(dim, s)] = adj
                self.flags[(dim, s)] = flg
                self.scores[(dim, s)] = sc

    # ---------- 07 ----------
    def aggregate(self):
        p = self.p
        for dim in DIM_ORDER:
            inc = p.dims_incluir.get(dim, 0)
            rows = [r for r in self.base_indice[dim] if r is not None]
            n_used = len(rows)
            cob_min = 0 if n_used == 0 else math.ceil(n_used * p.cobertura_sub - 1e-12)
            out = []
            for t in range(NY):
                if inc != 1:
                    out.append(None); continue
                num = den = 0.0; cnt = 0
                for s in range(1, 11):
                    r = self.base_indice[dim][s - 1]
                    if r is None:
                        continue
                    w = float(r.peso) if _num(r.peso) else 0.0
                    sc = self.scores[(dim, s)][t]
                    if sc is not None:
                        cnt += 1; num += sc * w; den += w
                if cnt < cob_min or den == 0:
                    out.append(None)
                else:
                    out.append(num / den)
            self.subindices[dim] = out
        for t in range(NY):
            lnsum = wsum = 0.0; cnt = 0; nonpos = 0
            for dim in DIM_ORDER:
                if p.dims_incluir.get(dim, 0) != 1:
                    continue
                s = self.subindices[dim][t]
                if s is None:
                    continue
                w = float(p.dims_peso.get(dim, 1))
                cnt += 1
                if s <= p.esc_min:
                    nonpos += 1
                else:
                    lnsum += math.log(s) * w; wsum += w
            if cnt < p.min_subindices:
                self.igda[t] = None
            elif nonpos > 0:
                self.igda[t] = p.esc_min
            else:
                self.igda[t] = math.exp(lnsum / wsum)
        for dim in DIM_ORDER:
            s0 = self.subindices[dim][0]
            self.escala[dim] = [None if (s0 in (None, 0) or v is None) else excel_round(100 * v / s0, 1) for v in self.subindices[dim]]
        i0 = self.igda[0]
        self.escala["IGDA"] = [None if (i0 in (None, 0) or v is None) else excel_round(100 * v / i0, 1) for v in self.igda]

    def run(self):
        self.score_rows(); self.build_index(); self.aggregate()
        return self

    # ---------- helpers ----------
    def selected(self):
        return [r for r in self.rows if r.seleccionado == 1]

    def summary(self):
        lines = []
        for dim, short in zip(DIM_ORDER, DIM_SHORT):
            s = self.subindices[dim]
            if all(v is None for v in s):
                continue
            lines.append(f"{short:17}" + " ".join(f"{v:6.2f}" if v is not None else "   -  " for v in s))
        lines.append(f"{'IGDA':17}" + " ".join(f"{v:6.2f}" if v is not None else "   -  " for v in self.igda))
        return "\n".join(lines)
