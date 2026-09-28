"""
Constrói IGDA_BDA_Construtor_v8.xlsx a partir da v7, aplicando v8_plan.py:
  1. copia o livro v7 (fórmulas intactas)
  2. alarga os intervalos $5:$335 -> $5:$<nova última linha> (fórmulas, filtro automático, formatação condicional, validação de dados)
  3. actualiza séries, metas, metadados e notas nas linhas existentes; sincroniza 10_Fontes
  4. acrescenta novas linhas de candidatos (fórmulas traduzidas da linha-modelo)
  5. acrescenta coluna "Origem da Meta 2027" (AT) e preenche (metas directas, transpostas, convertidas, MINPLAN, operacionais)
  6. uniformiza a grafia (pré-Acordo) dos nomes de dimensões/indicadores; corrige textos herdados (09_Metodologia, 01_Criterios, 10_Fontes)
  7. acrescenta a navegação do painel (folhas 11-13) e corrige o título do gráfico do IGDA
  8. devolve o log de alterações, a lista de metas alteradas e a categoria de cada meta
"""
from __future__ import annotations
import os, re, sys
from copy import copy

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.hyperlink import Hyperlink

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from xlsx_tools import *  # noqa
import v8_plan as P

REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
SRC = os.path.join(REPO, "Versão Final/Finalíssima/files/IGDA_BDA_Construtor_v7.xlsx")
YEARS = list(range(2015, 2026))


def _fmt(v, vazio="vazio"):
    """Formata valores para o registo de alterações (None -> texto explícito)."""
    if v is None:
        return vazio
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return f"{v:g}"
    return str(v)


def _pt(x, nd=None):
    if x is None:
        return "–"
    if nd is None:
        s = f"{x:g}"
    else:
        s = f"{x:.{nd}f}"
    return s.replace(".", ",")


def transpor(spec: dict, series: dict, ref_year: int = 2022):
    """Aplica a variação do PDN ao valor da série (ver v8_plan.TRANSPOR). Devolve (meta, texto, ano_usado, valor_usado)."""
    v = series.get(ref_year)
    year = ref_year
    if v is None:
        cands = [(abs(y - ref_year), 0 if y <= ref_year else 1, y) for y, val in series.items() if val is not None]
        if not cands:
            return None, f"{spec['pagina']}: {spec['conceito']} {_pt(spec['base'])} (2022) → {_pt(spec['meta'])} (2027); a série não tem observações para transpor a meta", None, None
        _, _, year = min(cands)
        v = series[year]
    base, meta, nd = spec["base"], spec["meta"], spec.get("nd", 1)
    if spec["modo"] == "add":
        m = round(v + (meta - base), nd)
        var = f"variação aditiva do PDN ({'+' if meta - base >= 0 else '−'}{_pt(abs(meta - base))} p.p.)"
    else:
        m = round(v * meta / base, nd)
        var = f"variação relativa do PDN ({(meta / base - 1) * 100:+.1f}%)".replace(".", ",").replace("+", "+").replace("-", "−")
    if nd == 0:
        m = int(round(m))
    txt = (f"{spec['pagina']}: {spec['conceito']} {_pt(base)} ({ref_year}) → {_pt(meta)} (2027) na base do PDN. A série ({spec['fonte_serie']}) regista {_pt(v, 2 if nd else 1)} em {year}"
           f"{'' if year == ref_year else ' (observação mais próxima de 2022)'}, pelo que a meta é TRANSPOSTA à base da série pela {var}: {_pt(m)}.")
    if spec.get("extra"):
        txt += f" Nota: {spec['extra']}."
    return m, txt, year, v


def build(out_path: str, version_label="v8", data_label="Setembro de 2026"):
    log = []
    wb = openpyxl.load_workbook(SRC)
    ws = wb[BASE]
    cols = sheet_columns(ws)
    old_last = last_data_row(ws)

    series = dict(P.SERIES)
    rows_new = list(P.NEW_ROWS)
    metas = dict(P.METAS)
    rename = dict(P.RENAME)

    new_last = old_last + len(rows_new)
    n = extend_ranges(wb, old_last, new_last)
    log.append(dict(tipo="estrutura", id="", descricao=f"Intervalos da Base_Potencial alargados de $5:${old_last} para $5:${new_last} ({n} fórmulas actualizadas)"))

    # ---- coluna Origem da Meta 2027 (AT) ----
    meta_col = ws.max_column + 1
    hdr_src = ws.cell(HDR_ROW, ws.max_column)
    c = ws.cell(HDR_ROW, meta_col, "Origem da Meta 2027")
    copy_style(hdr_src, c)
    ws.column_dimensions[get_column_letter(meta_col)].width = 70
    for r in range(FIRST, new_last + 1):
        copy_style(ws.cell(r, ws.max_column - 1), ws.cell(r, meta_col))
    log.append(dict(tipo="estrutura", id="", descricao="Nova coluna 'Origem da Meta 2027' (AT) na Base_Potencial, documentando a fonte e a categoria de cada meta"))

    # ---- filtro automático, formatação condicional e validação de dados alargados ----
    ws.auto_filter.ref = f"A{HDR_ROW}:{get_column_letter(meta_col)}{new_last}"
    new_cf = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        sq = " ".join(re.sub(r"(\D)%d$" % old_last, r"\g<1>%d" % new_last, part) for part in str(cf.sqref).split())
        for rule in cf.rules:
            new_cf.add(sq, rule)
    ws.conditional_formatting = new_cf
    for dv in ws.data_validations.dataValidation:
        letters = sorted({re.match(r"([A-Z]+)", str(rng).split(":")[0]).group(1) for rng in str(dv.sqref).split()})
        dv.sqref = MultiCellRange(" ".join(f"{L}{FIRST}:{L}{new_last}" for L in letters))
    log.append(dict(tipo="estrutura", id="", descricao=f"Filtro automático (A4:{get_column_letter(meta_col)}{new_last}), formatação condicional (AK/AL/AQ) e lista de validação de 'Ajuste manual' (AP) alargados às linhas novas e à coluna AT"))

    # ---- séries / metadados existentes ----
    def read_series(r):
        return {y: ws.cell(r, cols[y]).value for y in YEARS}

    def apply_spec(ind_id, spec, default_tipo):
        r = find_row_by_id(ws, ind_id)
        before = read_series(r)
        changed = []
        if spec.get("values"):
            for y, v in spec["values"].items():
                ws.cell(r, cols[y]).value = v
            changed = [f"{y}: {_fmt(before[y])}→{_fmt(v)}" for y, v in spec["values"].items()
                       if (before[y] is None) != (v is None) or (v is not None and before[y] is not None and abs(float(before[y]) - float(v)) > 1e-9)]
        for k, colname in (("indicador", "Indicador"), ("codigo", "Código/API"), ("unidade", "Unidade"), ("fonte", "Fonte"), ("origem", "Origem")):
            if spec.get(k):
                ws.cell(r, cols[colname]).value = spec[k]
        if spec.get("nacional") is not None:
            ws.cell(r, cols["Fonte nacional"]).value = spec["nacional"]
        if spec.get("minimo") is not None:
            ws.cell(r, cols["Mínimo"]).value = spec["minimo"]
        if spec.get("maximo") is not None:
            ws.cell(r, cols["Máximo"]).value = spec["maximo"]
        nota = ws.cell(r, cols["Nota"]).value or ""
        for old, new in spec.get("nota_sub", []):
            nota = nota.replace(old, new)
        if spec.get("nota_add"):
            nota = (nota + " | " if nota else "") + spec["nota_add"]
        ws.cell(r, cols["Nota"]).value = nota
        tipo = spec.get("tipo", default_tipo)
        desc = (spec.get("nota_add") or "metadados actualizados") + (" Alterações: " + "; ".join(changed) if changed else (" (sem alteração de valores)" if tipo == "série" else ""))
        log.append(dict(tipo=tipo, id=ind_id, descricao=desc))

    for ind_id, spec in series.items():
        apply_spec(ind_id, spec, "série")
    for ind_id, spec in rename.items():
        apply_spec(ind_id, spec, "metadados")

    # ---- metas ----
    meta_changes = []   # (id, antes, depois, origem, categoria)
    meta_cat = {}
    def set_meta(ind_id, meta, origem, cat):
        r = find_row_by_id(ws, ind_id)
        before = ws.cell(r, cols["Meta 2027"]).value
        ws.cell(r, cols["Meta 2027"]).value = meta
        ws.cell(r, meta_col).value = f"[{P.META_CAT_LABEL[cat]}] {origem}"
        meta_cat[ind_id] = cat
        if (before != meta) and not (before is None and meta is None):
            meta_changes.append((ind_id, before, meta, origem, cat))
            log.append(dict(tipo="meta", id=ind_id, descricao=f"Meta 2027: {_fmt(before, 'sem meta')} → {_fmt(meta, 'sem meta')} [{P.META_CAT_LABEL[cat]}]. {origem}"))

    for ind_id, (meta, origem, cat) in metas.items():
        set_meta(ind_id, meta, origem, cat)
    transposicoes = {}
    complementos = []
    for ind_id, spec in P.TRANSPOR.items():
        if spec["modo"] == "complemento":
            complementos.append((ind_id, spec)); continue
        r = find_row_by_id(ws, ind_id)
        m, txt, year, v = transpor(spec, read_series(r))
        set_meta(ind_id, m, txt, "transposta")
        transposicoes[ind_id] = dict(meta=m, ano=year, valor=v, base=spec["base"], meta_pdn=spec["meta"], modo=spec["modo"], conceito=spec["conceito"])
    for ind_id, spec in complementos:
        ref = P.TRANSPOR[spec["ref"]]
        m_ref = transposicoes[spec["ref"]]["meta"]
        m = round(100 - m_ref, spec.get("nd", 1))
        txt = f"{spec['pagina']}: {spec['conceito']} derivada como complemento (100 − meta transposta de {spec['ref']}, {_pt(m_ref)}): {_pt(m)}."
        set_meta(ind_id, m, txt, "transposta")
        transposicoes[ind_id] = dict(meta=m, ano=None, valor=None, base=None, meta_pdn=None, modo="complemento", conceito=spec["conceito"])
    # metas não alteradas: documentar origem genérica
    for r in range(FIRST, old_last + 1):
        if ws.cell(r, meta_col).value is None:
            m = ws.cell(r, cols["Meta 2027"]).value
            cat = "operacional" if m is not None else "sem_meta"
            meta_cat[ws.cell(r, 1).value] = cat
            ws.cell(r, meta_col).value = f"[{P.META_CAT_LABEL[cat]}] " + ("Meta operacional da v7 mantida (sem meta comparável no PDN 2023-2027)" if m is not None else "Sem meta definida")

    # ---- novas linhas ----
    template_row = old_last
    for i, spec in enumerate(rows_new):
        rr = old_last + 1 + i
        data = {"A": spec["id"], "B": spec["dim"], "C": spec["subtema"], "D": spec["indicador"], "E": spec["unidade"], "F": spec["sentido"],
                "G": spec["minimo"], "H": spec["maximo"], "I": spec.get("meta"), "J": spec.get("peso", 1), "K": spec["fonte"], "L": spec.get("codigo"),
                "M": spec["origem"], "N": spec.get("nacional", 0), "O": spec["grupo"], "P": spec["prioridade"], "AP": "Auto", "AS": spec.get("nota", "")}
        append_candidate_row(ws, template_row, rr, data)
        for y, v in spec["values"].items():
            ws.cell(rr, cols[y]).value = v
        cat = spec.get("meta_cat", "sem_meta" if spec.get("meta") is None else "operacional")
        meta_cat[spec["id"]] = cat
        ws.cell(rr, meta_col).value = f"[{P.META_CAT_LABEL[cat]}] {spec.get('meta_origem', '')}"
        log.append(dict(tipo="novo indicador", id=spec["id"], descricao=f"{spec['indicador']} — {spec['origem']}. {spec.get('nota', '')}"))

    # ---- 10_Fontes: sincronizar linhas alteradas e acrescentar as novas ----
    wf = wb["10_Fontes"]
    fontes_rows = {wf.cell(r, 1).value: r for r in range(5, wf.max_row + 1) if wf.cell(r, 1).value}
    synced = []
    for ind_id in list(series) + list(rename):
        fr = fontes_rows.get(ind_id)
        if not fr:
            continue
        br = find_row_by_id(ws, ind_id)
        wf.cell(fr, 2).value = ws.cell(br, cols["Indicador"]).value
        wf.cell(fr, 4).value = ws.cell(br, cols["Fonte"]).value
        if ws.cell(br, cols["Código/API"]).value:
            wf.cell(fr, 5).value = ws.cell(br, cols["Código/API"]).value
        wf.cell(fr, 8).value = ws.cell(br, cols["Origem"]).value
        spec = series.get(ind_id) or rename.get(ind_id)
        if spec.get("url"):
            wf.cell(fr, 7).value = spec["url"]
        synced.append(ind_id)
    lf = wf.max_row
    while lf > 4 and wf.cell(lf, 1).value in (None, ""):
        lf -= 1
    for i, spec in enumerate(rows_new):
        rr = lf + 1 + i
        for col in range(1, 9):
            copy_style(wf.cell(lf, col), wf.cell(rr, col))
        wf.cell(rr, 1, spec["id"]); wf.cell(rr, 2, spec["indicador"]); wf.cell(rr, 3, spec["dim"]); wf.cell(rr, 4, spec["fonte"])
        wf.cell(rr, 5, spec.get("codigo")); wf.cell(rr, 6, "Sim" if spec.get("nacional") else None); wf.cell(rr, 7, spec.get("url") or None); wf.cell(rr, 8, spec["origem"])
    wf.auto_filter.ref = f"A4:H{lf + len(rows_new)}"
    wf["B3"].value = "Uma linha por indicador do catálogo. As fontes nacionais provêm do INE, BNA, MINFIN, IIMS, MINPLAN, INSS/MAPTSS e ministérios sectoriais (o marcador identifica o produtor primário dos dados, independentemente do canal de extracção)."
    log.append(dict(tipo="estrutura", id="", descricao=f"10_Fontes sincronizada com a Base_Potencial para as {len(synced)} linhas alteradas ({', '.join(synced)}); 7 linhas novas com URL; filtro A4:H{lf + len(rows_new)}"))

    # ---- 01_Criterios: definição do marcador de fonte nacional ----
    wc = wb["01_Criterios"]
    for r in range(5, wc.max_row + 1):
        if isinstance(wc.cell(r, 2).value, str) and wc.cell(r, 2).value.startswith("Bónus fonte nacional"):
            wc.cell(r, 4).value = "Acresce ao score de indicadores cujo produtor primário é nacional (INE, BNA, MINFIN, IIMS, ministérios sectoriais), independentemente do canal de extracção (WDI, WEO, ILOSTAT)."

    # ---- grafia (pré-Acordo) ----
    n_orto = fix_orthography(wb, ws, wf, cols)
    log.append(dict(tipo="estrutura", id="", descricao=f"Grafia uniformizada para a norma pré-Acordo Ortográfico em {n_orto} células (nome da dimensão 8 'Sector Privado' em 01_Criterios, 02_Base_Potencial, 03_Base_Indice e 10_Fontes; nomes de indicadores e unidades)"))

    # ---- 09_Metodologia: textos herdados ----
    w9 = wb["09_Metodologia"]
    n_cand = new_last - FIRST + 1
    for row in w9.iter_rows(min_row=1, max_row=w9.max_row):
        for cc in row:
            v = cc.value
            if not isinstance(v, str):
                continue
            if "331 candidatos" in v:
                v = v.replace("331 candidatos", f"{n_cand} candidatos (331 na v7 + {len(rows_new)} acrescentados na v8: LAB028-LAB031, INC031-INC033)")
                v = v.replace("IIMS 2023-24)", "IIMS 2023-24; INSS/MAPTSS; MINPLAN — Balanço do PDN)")
            if v.strip().startswith("Min-Max para a escala 0–100 com limites fixos por indicador"):
                v = ("  Min-Max para a escala 0–100 com limites fixos por indicador, definidos por referência a valores externos e normativos (escala oficial da fonte, limite natural da variável "
                     "ou tecto operacional documentado). Nenhum dos 41 indicadores do índice usa limites derivados da amplitude histórica (revisão v7); alguns candidatos fora do índice mantêm ainda "
                     "esses limites, assinalados na coluna 'Nota', para revisão se vierem a ser incorporados. Justificação linha a linha no documento anexo 'Justificação dos limites mínimo e máximo (v8)'.")
            if "ex.: PIB não petrolífero, cobertura 4G, protecção social INSS" in v:
                v = v.replace("ex.: PIB não petrolífero, cobertura 4G, protecção social INSS", "ex.: taxa líquida de matrícula no primário, cobertura 4G, IDE não petrolífero")
            cc.value = v

    # ---- 00_Painel: navegação para as folhas 11-13 ----
    add_painel_navigation(wb)
    log.append(dict(tipo="estrutura", id="", descricao="00_Painel: navegação alargada às folhas 11_Escala_Comum, 12_Emprego_INE e 13_Alteracoes_v8"))

    # ---- 07_Subindices: título do gráfico do IGDA (herdado como 'None') ----
    for ch in wb["07_Subindices"]._charts:
        if (ch.anchor._from.col, ch.anchor._from.row) == (21, 14):
            ch.title = "IGDA-BDA (média geométrica), 2015–2025"
            log.append(dict(tipo="estrutura", id="", descricao="07_Subindices: título do gráfico do IGDA corrigido (a v7 mostrava um título inválido, herdado de uma célula vazia)"))

    # ---- textos de versão ----
    ws0 = wb["00_Painel"]
    ws0["B2"].value = f"ÍNDICE GLOBAL DE DESENVOLVIMENTO DE ANGOLA (IGDA-BDA) — Selecção Final · {version_label}"
    ws["B3"].value = (f"{n_cand} indicadores candidatos, valores anuais como extraídos das fontes de origem (INE, BNA, MINFIN/FMI, IIMS, MINPLAN, INSS/MAPTSS, Banco Mundial/WDI, "
                      "OIT, OMS, UNCTAD, benchmarks internacionais). Colunas de avaliação e selecção calculadas por fórmulas. Editável: Peso (J), Meta 2027 (I) e Ajuste manual (AP). "
                      f"A coluna 'Origem da Meta 2027' documenta a categoria e a fonte de cada meta ({version_label}).")
    return dict(wb=wb, log=log, new_last=new_last, meta_col=meta_col, meta_changes=meta_changes, meta_cat=meta_cat, transposicoes=transposicoes, n_cand=n_cand)


def fix_orthography(wb, ws, wf, cols):
    n = 0
    # 1) chave da dimensão 8 em todas as folhas (valor exacto ou contido em texto não-fórmula)
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("=") and P.DIM8_OLD in c.value:
                    c.value = c.value.replace(P.DIM8_OLD, P.DIM8_NEW); n += 1
    # 2) nomes de subtema, indicador e unidade (Base) e espelho em 10_Fontes
    def fix(cell):
        nonlocal n
        v = cell.value
        if not isinstance(v, str) or v.startswith("="):
            return
        nv = v
        for old, new in P.ORTOGRAFIA:
            nv = nv.replace(old, new)
        if nv != v:
            cell.value = nv; n += 1
    for r in range(FIRST, ws.max_row + 1):
        for colname in ("Subtema", "Indicador", "Unidade"):
            fix(ws.cell(r, cols[colname]))
    for r in range(5, wf.max_row + 1):
        fix(wf.cell(r, 2))
    return n


def add_painel_navigation(wb):
    ws = wb["00_Painel"]
    ins_at, k = 25, 3
    # objectos que o insert_rows não desloca: células unidas, formatação condicional, âncoras de gráficos, alturas de linha
    merged = [str(m) for m in ws.merged_cells.ranges]
    cfs = [(str(cf.sqref), list(cf.rules)) for cf in ws.conditional_formatting]
    heights = {r: ws.row_dimensions[r].height for r in list(ws.row_dimensions.keys()) if r >= ins_at}
    ws.insert_rows(ins_at, k)
    # insert_rows já deslocou as células (incluindo os marcadores MergedCell); só falta actualizar as coordenadas dos intervalos unidos
    from openpyxl.worksheet.cell_range import CellRange
    new_ranges = []
    for m in merged:
        cr = CellRange(m)
        if cr.min_row >= ins_at:
            cr.shift(row_shift=k)
        new_ranges.append(cr)
    ws.merged_cells = MultiCellRange(new_ranges)
    new_cf = ConditionalFormattingList()
    for sq, rules in cfs:
        parts = []
        for part in sq.split():
            cr = CellRange(part)
            if cr.min_row >= ins_at:
                cr.shift(row_shift=k)
            parts.append(cr.coord)
        for rule in rules:
            new_cf.add(" ".join(parts), rule)
    ws.conditional_formatting = new_cf
    for ch in ws._charts:
        a = ch.anchor
        if a._from.row + 1 >= ins_at:
            a._from.row += k
        if getattr(a, "to", None) is not None and a.to.row + 1 >= ins_at:
            a.to.row += k
    for r in sorted(heights, reverse=True):
        ws.row_dimensions[r + k].height = heights[r]
    for r in range(ins_at, ins_at + k):
        ws.row_dimensions[r].height = ws.row_dimensions[ins_at - 1].height
    links = [("→ 11 · Escala comum", "11_Escala_Comum", "Evolução dos subíndices e do IGDA em base 2015 = 100"),
             ("→ 12 · Emprego INE", "12_Emprego_INE", "Leitura complementar do emprego com séries INE/IEA 2019–2025 (nova na v8)"),
             ("→ 13 · Alterações v8", "13_Alteracoes_v8", "Registo linha a linha das alterações v7 → v8 e comparação de resultados (nova na v8)")]
    src_b, src_e = ws.cell(ins_at - 1, 2), ws.cell(ins_at - 1, 5)
    for i, (txt, sheet, desc) in enumerate(links):
        cb = ws.cell(ins_at + i, 2, txt); copy_style(src_b, cb)
        cb.hyperlink = Hyperlink(ref=cb.coordinate, location=f"'{sheet}'!A1", display=txt)
        ce = ws.cell(ins_at + i, 5, desc); copy_style(src_e, ce)
    # sanidade: nenhuma fórmula do livro referencia 00_Painel
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "00_Painel" in c.value:
                    raise RuntimeError(f"fórmula referencia 00_Painel: {sh.title}!{c.coordinate}")


if __name__ == "__main__":
    res = build("/tmp/test_v8.xlsx")
    res["wb"].save(os.path.join(HERE, "test_v8.xlsx"))
    for e in res["log"]:
        print(e["tipo"], e["id"], e["descricao"][:160])
    print("new_last", res["new_last"], "| metas alteradas", len(res["meta_changes"]))
    for t in res["meta_changes"]:
        print(t[0], t[1], "→", t[2], t[4])
