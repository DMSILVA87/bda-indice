"""
Utilitários python-pptx para actualizar a apresentação CEX preservando formatação:
  - set_text_preserve(shape, lines): substitui o texto de cada parágrafo mantendo o formato do 1º run
  - set_table_cell(table, r, c, text): substitui texto de célula preservando formato
  - replace_chart_series(chart, categories, series_dict)
  - find_shape(slide, name=None, contains=None)
"""
from __future__ import annotations
import copy
from pptx.chart.data import CategoryChartData


def _set_paragraph_text(paragraph, text):
    runs = paragraph.runs
    if not runs:
        paragraph.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_text_preserve(shape, lines):
    """lines: lista de strings, uma por parágrafo. Se houver mais linhas que parágrafos, clona o último parágrafo."""
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    # remover parágrafos a mais
    while len(paras) > len(lines):
        p = paras.pop()
        p._p.getparent().remove(p._p)
    # adicionar clones se faltam
    while len(paras) < len(lines):
        src = paras[-1]._p
        new = copy.deepcopy(src)
        src.addnext(new)
        paras = list(tf.paragraphs)
    for p, line in zip(tf.paragraphs, lines):
        _set_paragraph_text(p, line)


def set_table_cell(table, r, c, text):
    cell = table.cell(r, c)
    tf = cell.text_frame
    paras = list(tf.paragraphs)
    _set_paragraph_text(paras[0], text)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def find_shape(slide, name=None, contains=None, shape_id=None):
    for sh in slide.shapes:
        if shape_id is not None and sh.shape_id == shape_id:
            return sh
        if name is not None and sh.name == name:
            return sh
        if contains is not None and sh.has_text_frame and contains in sh.text_frame.text:
            return sh
    return None


def replace_chart_series(chart, categories, series: dict, number_format="0.0"):
    """Tenta substituir os dados (workbook embutido); se o workbook estiver ligado externamente, edita a cache XML das séries."""
    cd = CategoryChartData()
    cd.categories = list(categories)
    for name, vals in series.items():
        cd.add_series(name, list(vals), number_format=number_format)
    try:
        chart.replace_data(cd)
        return "replace_data"
    except ValueError:
        pass
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    plot = chart.plots[0]
    ser_els = plot._element.findall("c:ser", ns)
    vals_list = list(series.values())
    names = list(series.keys())
    for k, ser in enumerate(ser_els):
        if k >= len(vals_list):
            break
        cache = ser.find("c:val/c:numRef/c:numCache", ns)
        if cache is None:
            cache = ser.find("c:val/c:numLit", ns)
        if cache is not None:
            for pt in cache.findall("c:pt", ns):
                cache.remove(pt)
            ptc = cache.find("c:ptCount", ns)
            if ptc is not None:
                ptc.set("val", str(len(vals_list[k])))
            from lxml import etree
            for idx, v in enumerate(vals_list[k]):
                if v is None:
                    continue
                pt = etree.SubElement(cache, "{%s}pt" % ns["c"]); pt.set("idx", str(idx))
                ve = etree.SubElement(pt, "{%s}v" % ns["c"]); ve.text = repr(float(v))
        # nome da série (cache de texto)
        tx = ser.find("c:tx/c:strRef/c:strCache/c:pt/c:v", ns)
        if tx is not None and k < len(names):
            tx.text = names[k]
        # categorias
        cat = ser.find("c:cat/c:strRef/c:strCache", ns)
        if cat is None:
            cat = ser.find("c:cat/c:numRef/c:numCache", ns)
        if cat is not None:
            pts = cat.findall("c:pt", ns)
            for idx, pt in enumerate(pts):
                if idx < len(categories):
                    pt.find("c:v", ns).text = str(categories[idx])
    return "xml_cache"


def dump_text(shape):
    return [p.text for p in shape.text_frame.paragraphs] if shape.has_text_frame else None
