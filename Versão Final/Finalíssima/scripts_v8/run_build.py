import sys, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib, build_v8, build_v8_sheets, v8_plan, igda_engine
importlib.reload(v8_plan); importlib.reload(build_v8); importlib.reload(build_v8_sheets)
from igda_engine import Model, load_inputs_from_workbook, DIM_ORDER, DIM_SHORT

OUT = sys.argv[1] if len(sys.argv) > 1 else "test_v8.xlsx"
wb, log, new_last, meta_col = build_v8.build(OUT)
labour = json.load(open("labour_data.json"))
web_new = {"2026T1": {"desemprego": 21.3, "emprego": 42.3, "desemprego_jovem": 40.7, "informal": 79.3, "actividade": 53.8},
           "2026T2": {"desemprego": 21.5, "emprego": 44.2, "desemprego_jovem": 40.5, "informal": 80.1, "actividade": 56.3}}
build_v8_sheets.add_emprego_ine(wb, labour, web_new)
# results v7 (recalculated) and v8 (engine on the new workbook, saved first to a temp)
tmp = OUT.replace(".xlsx", "_pre.xlsx"); wb.save(tmp)
m7 = Model(load_inputs_from_workbook("/home/user/bda-indice/Versão Final/Finalíssima/files/IGDA_BDA_Construtor_v7.xlsx")).run()
m8 = Model(load_inputs_from_workbook(tmp)).run()
def pack(m):
    dims = {s: m.subindices[d] for d, s in zip(DIM_ORDER, DIM_SHORT) if any(v is not None for v in m.subindices[d])}
    return {"dims": dims, "igda": m.igda}
# meta changes list
import openpyxl
v7 = openpyxl.load_workbook("/home/user/bda-indice/Versão Final/Finalíssima/files/IGDA_BDA_Construtor_v7.xlsx", data_only=True)["02_Base_Potencial"]
v7meta = {r[0]: r[8] for r in v7.iter_rows(min_row=5, values_only=True) if r[0]}
meta_changes = []
for mid, (meta, origem) in v8_plan.METAS.items():
    if mid in v7meta and (v7meta[mid] != meta):
        meta_changes.append((mid, v7meta[mid], meta, origem))
build_v8_sheets.add_changelog(wb, log, pack(m7), pack(m8), meta_changes)
build_v8_sheets.add_version_notes(wb)
wb.save(OUT)
json.dump({"log": log, "meta_changes": meta_changes, "v7": pack(m7), "v8": pack(m8), "selected_v8": [r.id for r in m8.selected()]}, open(OUT.replace(".xlsx", "_results.json"), "w"), ensure_ascii=False, indent=1, default=str)
print("saved", OUT, "| rows", len(m8.rows), "| selected", len(m8.selected()), "| sheets", wb.sheetnames)
print(m8.summary())
